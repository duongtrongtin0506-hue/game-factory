"""Ảnh xem trước THẬT cho mẫu .xlsx: tính lại toàn bộ công thức bằng bộ tính độc lập (formulas), dựng HTML giống bảng tính
từ giá trị đã tính, chụp bằng Chrome headless (xuong/chup.py). Không vẽ giả: số trên ảnh = số trong tệp.
  python xem_truoc.py <file.xlsx> <thư-mục-ra> "Sheet:hàng:cột:ten-anh" ...
"""
import datetime as dt
import html
import sys
from pathlib import Path

import formulas
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import kiem_cong_thuc  # noqa: E402,F401  (đăng ký SUMIFS/COUNTIFS/REPT cho bộ tính)
from chup import shot  # noqa: E402

CSS = """body{margin:0;background:#fff;font:15px/1.4 Arial,Helvetica,sans-serif;color:#1D1D1F}
.bar{display:flex;gap:2px;background:#F5F5F7;border-bottom:1px solid #D2D2D7;padding:8px 8px 0;font-size:13px}.bar span{padding:6px 14px;border-radius:8px 8px 0 0;color:#6E6E73}.bar span.on{background:#fff;color:#1D1D1F;font-weight:700;border:1px solid #D2D2D7;border-bottom:0}
table{border-collapse:collapse;margin:0}td,th{border:1px solid #ECECF0;padding:6px 8px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;max-width:320px;height:22px}
th{background:#F5F5F7;color:#8E8E93;font-weight:400;font-size:12px;text-align:center}td.rn{background:#F5F5F7;color:#8E8E93;font-size:12px;text-align:center;width:28px}td.n{text-align:right;font-variant-numeric:tabular-nums}"""


def values(path):
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    out = {}
    for k, rng in sol.items():
        try:
            v = rng.value.ravel().tolist()
        except Exception:
            continue
        ref = k.split("]")[-1]
        sh, _, cell = ref.rpartition("!")
        sh = sh.strip("'").upper()
        if ":" in cell or len(v) != 1:
            continue
        out[(sh, cell.replace("$", "").upper())] = v[0]
    return out


def fmt(v, nf):
    if v is None or v == "":
        return ""
    if isinstance(v, (dt.date, dt.datetime)):
        v = (v - dt.datetime(1899, 12, 30)).days if isinstance(v, dt.datetime) else (v - dt.date(1899, 12, 30)).days
    if isinstance(v, bool):
        return str(v).upper()
    if isinstance(v, (int, float)):
        nf = nf or "General"
        if "$" in nf:
            return f"${v:,.2f}" if ".00" in nf else f"${v:,.0f}"
        if "%" in nf:
            return f"{v * 100:.0f}%"
        if any(t in nf for t in ("yy", "mmm", "d")) and 20000 < v < 80000:
            d = dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
            if "ddd" in nf: return d.strftime("%a %b %-d") if sys.platform != "win32" else d.strftime("%a %b %#d")
            if "d" in nf and "yyyy" in nf: return d.strftime("%b %#d, %Y") if sys.platform == "win32" else d.strftime("%b %-d, %Y")
            if "yyyy" in nf: return d.strftime("%b %Y")
            return d.strftime("%b")
        return f"{v:,.0f}" if float(v).is_integer() else f"{v:,.2f}"
    s = str(v)
    return "" if s.startswith("#") and s.endswith(("!", "?", "A")) else s


def render(path, vals, sheet, rows, cols, out_png, w=1280, h=800):
    wb = load_workbook(path)
    ws = wb[sheet]
    tabs = "".join(f'<span class="{"on" if s.title == sheet else ""}">{html.escape(s.title)}</span>' for s in wb.worksheets if s.sheet_state == "visible")
    from openpyxl.utils import get_column_letter as L
    head = "<tr><th></th>" + "".join(f'<th style="width:{(ws.column_dimensions[L(c)].width or 10) * 7.5:.0f}px">{L(c)}</th>' for c in range(1, cols + 1)) + "</tr>"
    body = []
    for r in range(1, rows + 1):
        tds = [f'<td class="rn">{r}</td>']
        for c in range(1, cols + 1):
            cell = ws.cell(r, c)
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                v = vals.get((sheet.upper(), cell.coordinate))
            txt = fmt(v, cell.number_format)
            st = []
            f = cell.font
            if f is not None:
                if f.b: st.append("font-weight:700")
                if f.sz and f.sz != 11: st.append(f"font-size:{min(f.sz, 22) * 1.33:.0f}px")
                if f.color is not None and isinstance(f.color.rgb, str) and len(f.color.rgb) >= 6: st.append(f"color:#{f.color.rgb[-6:]}")
            if cell.fill is not None and cell.fill.fill_type == "solid" and isinstance(cell.fill.fgColor.rgb, str):
                st.append(f"background:#{cell.fill.fgColor.rgb[-6:]}")
            cls = ' class="n"' if isinstance(v, (int, float)) and not isinstance(v, bool) else ""
            tds.append(f'<td{cls} style="{";".join(st)}">{html.escape(txt)}</td>')
        body.append("<tr>" + "".join(tds) + "</tr>")
    page = f'<!doctype html><meta charset="utf-8"><style>{CSS}</style><div class="bar">{tabs}</div><table>{head}{"".join(body)}</table>'
    tmp = Path(out_png).with_suffix(".html")
    tmp.write_text(page, encoding="utf-8")
    ok = shot(str(tmp), str(out_png), f"{w}x{h}")
    if not __import__("os").environ.get("N24_GIU_HTML"):
        tmp.unlink(missing_ok=True)
    return ok


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    p, out = Path(sys.argv[1]), Path(sys.argv[2])
    out.mkdir(parents=True, exist_ok=True)
    vals = values(p)
    for spec in sys.argv[3:]:
        sh, rows, cols, name = spec.split(":")
        print(name, render(p, vals, sh, int(rows), int(cols), out / f"{name}.png"))
