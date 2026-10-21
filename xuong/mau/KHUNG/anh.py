"""Preview images for a template, drawn from the workbook's own computed values (not a mock-up):
python KHUNG/anh.py <file.xlsx> <out-dir> "<Title>" "<tagline>"
Writes preview-sheet-1600x1200.png (main sheet as it looks, values from the formulas engine) and
cover-1600x1200.png (title + the sheet image), via Chromium (Playwright).
"""
import html
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import kiem_cong_thuc  # noqa: E402,F401  (registers the *IFS functions)
import formulas  # noqa: E402

ACCENT = "#E8833A"


def fmt(v, nf):
    if v is None or str(v) in ("empty", "None"):
        return ""
    try:
        f = float(v)
    except Exception:
        return html.escape(str(v))
    if "%" in nf:
        return f"{f * 100:.0f}%"
    if "$" in nf:
        return f"${f:,.0f}" if "0.00" not in nf else f"${f:,.2f}"
    if "yy" in nf or "dd" in nf or "mmm" in nf:
        import datetime
        d = datetime.date(1899, 12, 30) + datetime.timedelta(days=int(f))
        return d.strftime("%a %d %b") if "ddd" in nf else d.isoformat()
    if "0.0" in nf:
        return f"{f:.1f}"
    return f"{f:,.0f}" if abs(f) >= 1000 or f == int(f) else f"{f:,.2f}"


def sheet_html(path, sheet, max_rows=34, max_cols=14):
    wb = load_workbook(path)
    ws = wb[sheet]
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    vals = {}
    for key, rng in sol.items():
        k = key.split("]")[-1].replace("'", "")
        if "!" in k and ":" not in k:
            try:
                vals[k.upper()] = rng.value.ravel()[0]
            except Exception:
                pass
    rows = []
    for r in range(1, min(ws.max_row, max_rows) + 1):
        cells = []
        for c in range(1, min(ws.max_column, max_cols) + 1):
            if ws.column_dimensions[get_column_letter(c)].hidden:
                continue
            cell = ws.cell(row=r, column=c)
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                v = vals.get(f"{sheet.upper()}!{get_column_letter(c)}{r}", "")
            elif hasattr(v, "isoformat"):
                import datetime
                v = (v.date() if isinstance(v, datetime.datetime) else v) - datetime.date(1899, 12, 30)
                v = v.days
            txt = fmt(v, cell.number_format or "")
            bg = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
            style = []
            if bg and isinstance(bg, str) and len(bg) >= 6:
                style.append(f"background:#{bg[-6:]}")
            if cell.font and cell.font.color and isinstance(cell.font.color.rgb, str) and cell.font.color.rgb[-6:] == "FFFFFF":
                style.append("color:#fff;font-weight:600")
            if r == 1:
                style.append("font-size:26px;font-family:Georgia,serif")
            w = ws.column_dimensions[get_column_letter(c)].width or 9
            cells.append(f'<td style="{";".join(style)};min-width:{w * 7.2:.0f}px">{txt}</td>')
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return "<table>" + "".join(rows) + "</table>"


def render(path, out, title, tagline):
    from playwright.sync_api import sync_playwright
    wb = load_workbook(path)
    main = wb.sheetnames[1]
    table = sheet_html(path, main)
    css = ("body{margin:0;background:#FFFBF5;font-family:Inter,Segoe UI,Arial,sans-serif;color:#1A1A1A}"
           "table{border-collapse:collapse;font-size:13px}td{border-bottom:1px solid #E8E0D5;padding:5px 8px;white-space:nowrap;height:18px}"
           ".sheet{background:#fff;border-radius:16px;box-shadow:0 2px 24px rgba(0,0,0,.08);padding:24px;overflow:hidden}"
           ".tabs{display:flex;gap:6px;margin:10px 0 0 8px}.tabs span{font-size:12px;padding:6px 12px;border-radius:8px 8px 0 0;background:#EFE7DC}.tabs span.on{background:#fff;font-weight:600;color:" + ACCENT + "}")
    tabs = "".join(f'<span class="{"on" if n == main else ""}">{html.escape(n)}</span>' for n in wb.sheetnames)
    sheet_page = f"<html><head><meta charset='utf-8'><style>{css}</style></head><body style='padding:40px;width:1520px;height:1120px;overflow:hidden'><div class='sheet'>{table}</div><div class='tabs'>{tabs}</div></body></html>"
    cover = (f"<html><head><meta charset='utf-8'><style>{css}"
             ".wrap{display:grid;grid-template-columns:520px 1fr;height:1200px}.left{padding:96px 56px;display:flex;flex-direction:column;justify-content:center}"
             ".lab{font:600 13px/1 ui-monospace,Consolas,monospace;letter-spacing:.12em;color:#8B7E6F;text-transform:uppercase}"
             f"h1{{font:400 64px/1 Georgia,'Palatino Linotype',serif;margin:24px 0}}.pill{{display:inline-block;background:{ACCENT};color:#fff;border-radius:999px;padding:14px 28px;font-weight:600;margin-top:24px}}"
             ".right{padding:72px 0 72px 0;overflow:hidden}.right .sheet{transform:scale(.82);transform-origin:top left;width:1300px}</style></head>"
             f"<body><div class='wrap'><div class='left'><div class='lab'>Free template · Google Sheets + Excel</div><h1>{html.escape(title)}</h1>"
             f"<div style='font-size:20px;color:#5A5048'>{html.escape(tagline)}</div><div><span class='pill'>Get it free</span></div></div>"
             f"<div class='right'><div class='sheet'>{table}</div></div></div></body></html>")
    out.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as pw:
        b = pw.chromium.launch()
        p = b.new_page(viewport={"width": 1600, "height": 1200})
        for name, page in (("preview-sheet-1600x1200.png", sheet_page), ("cover-1600x1200.png", cover)):
            p.set_content(page)
            p.wait_for_timeout(200)
            p.screenshot(path=str(out / name))
        b.close()


if __name__ == "__main__":
    render(Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3], sys.argv[4])
