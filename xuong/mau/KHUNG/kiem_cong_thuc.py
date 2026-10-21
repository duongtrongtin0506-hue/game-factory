"""Formula check (Cửa 0 for templates): python KHUNG/kiem_cong_thuc.py <file.xlsx> [Sheet!A1 ...]

Recalculates the whole workbook with the `formulas` engine (independent of Excel/Sheets) and fails if any
formula cell evaluates to an error (#VALUE!, #NAME?, #REF!, #DIV/0!, #N/A, #NUM!). Prints the listed cells.
Writes <file>.kiem.json next to the workbook: {"formulas": n, "errors": [...], "values": {...}}.
"""
import json
import sys
from pathlib import Path

import formulas
from openpyxl import load_workbook

ERR = ("#VALUE!", "#NAME?", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")

# The engine lacks the *IFS family; add them with Excel's criteria rules (=, <>, <, <=, >, >=, blank, non-blank).
import datetime as _dt
import re as _re

import numpy as _np


def _num(v):
    if isinstance(v, (int, float, _np.integer, _np.floating)) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, (_dt.date, _dt.datetime)):
        return float((v - _dt.date(1899, 12, 30)).days) if isinstance(v, _dt.date) else None
    try:
        return float(str(v))
    except Exception:
        return None


def _blank(v):
    return v is None or str(v) in ("", "empty") or type(v).__name__ == "XlError" and False or str(v) == "None"


def _match(v, crit):
    c = getattr(crit, 'value', crit)
    if hasattr(c, "ravel"):
        c = c.ravel()[0]
    if isinstance(c, (int, float, _np.integer, _np.floating)):
        n = _num(v)
        return n is not None and n == float(c)
    c = str(c)
    m = _re.match(r"^(<=|>=|<>|<|>|=)?(.*)$", c, _re.S)
    op, rhs = m.group(1) or "=", m.group(2)
    blank = _blank(v) or str(v) == "empty"
    if rhs == "":
        return (not blank) if op == "<>" else blank if op == "=" else False
    rn, vn = _num(rhs), (None if blank else _num(v))
    if rn is not None and vn is not None and op in ("<", "<=", ">", ">=", "=", "<>"):
        return {"<": vn < rn, "<=": vn <= rn, ">": vn > rn, ">=": vn >= rn, "=": vn == rn, "<>": vn != rn}[op]
    sv = "" if blank else str(v).lower()
    if op == "=":
        return sv == rhs.lower()
    if op == "<>":
        return sv != rhs.lower()
    return False


def _flat(a):
    a = getattr(a, 'value', a)
    return list(_np.asarray(a, dtype=object).ravel())


def _mask(pairs):
    mask = None
    for rng, crit in pairs:
        vals = _flat(rng)
        m = [_match(v, crit) for v in vals]
        mask = m if mask is None else [a and b for a, b in zip(mask, m)]
    return mask


def _sumifs(sum_range, *args):
    pairs = list(zip(args[0::2], args[1::2]))
    vals = _flat(sum_range)
    return sum((_num(v) or 0.0) for v, ok in zip(vals, _mask(pairs)) if ok)


def _countifs(*args):
    return float(sum(1 for ok in _mask(list(zip(args[0::2], args[1::2]))) if ok))


def _averageifs(avg_range, *args):
    pairs = list(zip(args[0::2], args[1::2]))
    xs = [(_num(v) or 0.0) for v, ok in zip(_flat(avg_range), _mask(pairs)) if ok and _num(v) is not None]
    return sum(xs) / len(xs) if xs else float("nan")


def _n(x):
    v = _np.asarray(getattr(x, "value", x), dtype=object)
    f = _np.vectorize(lambda e: float(e) if isinstance(e, (int, float, _np.integer, _np.floating)) and not isinstance(e, bool) else 0.0, otypes=[float])
    out = f(v)
    return out if out.shape else float(out)


_F = formulas.get_functions()
_F["N"] = _n
_F["SUMIFS"] = _sumifs
_F["COUNTIFS"] = _countifs
_F["AVERAGEIFS"] = _averageifs


def check(path, cells=()):
    path = Path(path)
    wb = load_workbook(path)
    n = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    errors, values = [], {}
    for key, rng in sol.items():
        try:
            vals = rng.value.ravel().tolist()
        except Exception:
            continue
        for v in vals:
            s = str(v)
            if any(e in s for e in ERR):
                errors.append(f"{key}: {s}")
                break
    want = {c.upper() for c in cells}
    for key, rng in sol.items():
        k = key.split("]")[-1].replace("'", "").upper()
        if k in want:
            try:
                values[k] = rng.value.ravel().tolist()[0]
            except Exception:
                values[k] = str(rng.value)
    res = {"formulas": n, "errors": errors[:30], "error_count": len(errors), "values": {k: (v if isinstance(v, (int, float, str)) else str(v)) for k, v in values.items()}}
    Path(str(path) + ".kiem.json").write_text(json.dumps(res, indent=2, ensure_ascii=False), encoding="utf-8")
    return res


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    r = check(sys.argv[1], sys.argv[2:])
    print(json.dumps(r, ensure_ascii=False)[:1500])
    sys.exit(1 if r["error_count"] else 0)
